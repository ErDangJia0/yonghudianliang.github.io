# -*- coding: utf-8 -*-
"""流式读取电量表，自动识别两种结构：
  模式 A (detailed)：分时电量 sheet（96 个 HH:MM 列）+ 分时电价 sheet（24 小时价）
  模式 B (settlement)：结算汇总 sheet（月份×企业×结算电量×结算电价）
"""
import json
import os
import sys
import time
import openpyxl

# ---------- 自动探测源文件（跳过 Excel 锁文件 ~$xxx.xlsx） ----------
def _list_xlsx(top):
    try:
        for f in os.listdir(top):
            if f.endswith(".xlsx") and not f.startswith("~$"):
                yield os.path.join(top, f)
    except FileNotFoundError:
        pass

# 支持 python generate_data.py "C:\xxx\新表.xlsx" 手动指定
SRC = sys.argv[1] if len(sys.argv) > 1 else None
if SRC and not os.path.exists(SRC):
    SRC = None

if SRC is None:
    CANDIDATES = list(_list_xlsx(r"C:\Users\WHY\Desktop")) + list(_list_xlsx(os.path.dirname(os.path.abspath(__file__))))
    best_mtime = 0
    for p in CANDIDATES:
        try:
            mt = os.path.getmtime(p)
            if mt > best_mtime:
                best_mtime = mt
                SRC = p
        except FileNotFoundError:
            pass
if not SRC:
    raise SystemExit("未找到电量表 xlsx")
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data.js")
print(f"[source] {SRC} ({os.path.getsize(SRC)/1024/1024:.1f} MB)")
t0 = time.time()

wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
print(f"[openpyxl] {time.time()-t0:.1f}s, sheets: {wb.sheetnames}")


def _read_header(ws):
    rows = ws.iter_rows(values_only=True)
    h = next(rows)
    return [str(x) if x is not None else "" for x in h], rows


def _detect_time_cols(header):
    cols = []
    for i, h in enumerate(header):
        if h == "24:00" or (":" in h and len(h.split(":")) == 2):
            try:
                hh, mm = h.split(":")
                cols.append((i, int(hh), int(mm), h))
            except Exception:
                pass
    cols.sort(key=lambda t: (t[1] * 60 + t[2]))
    return cols


def slot_hour(hh, mm):
    """label 'HH:MM' 是时段结束时间，对应 hour = floor((HH:MM - 15min) / 1h)"""
    total_min = hh * 60 + mm - 15
    return max(0, min(23, total_min // 60))


def month_key(x):
    s = str(x or "").replace("月", "").strip()
    if not s:
        return None
    # 尝试 YYYY-MM 格式
    if "-" in s and len(s.split("-")[1]) == 2:
        return s if len(s) == 7 else f"2026-{int(s.split('-')[1]):02d}"
    try:
        return f"2026-{int(s):02d}"
    except Exception:
        return s


# ================================================================
# 识别模式
# ================================================================
mode = None
for sn in wb.sheetnames:
    h, _ = _read_header(wb[sn])
    tc = _detect_time_cols(h)
    if len(tc) >= 90:
        mode = "detailed"
        break

if mode is None:
    # 结算汇总模式：找含"结算电量"/"结算电价"的 sheet
    for sn in wb.sheetnames:
        h, _ = _read_header(wb[sn])
        names = [x for x in h if x]
        if any("结算电量" in x or ("电量" in x and "兆瓦" in x) for x in names) and \
           any("结算电价" in x or ("电价" in x and "兆瓦" in x) for x in names):
            mode = "settlement"
            break

if mode is None:
    # 回退：根据 sheet 数量猜测
    mode = "detailed" if len(wb.sheetnames) >= 2 else "settlement"
print(f"[mode] {mode}")


# ================================================================
# 模式 A：分时
# ================================================================
if mode == "detailed":
    # ---- 电价 ----
    px_sheet = None
    for sn in wb.sheetnames:
        if sn != wb.sheetnames[0]:
            px_sheet = wb[sn]
            break
    if px_sheet is None:
        px_sheet = wb[wb.sheetnames[0]]
    price_map = {}
    h, rows = _read_header(px_sheet)
    idx_month = next((i for i, x in enumerate(h) if "月" in x), 0)
    idx_comp = next((i for i, x in enumerate(h) if "企业" in x), 1)
    # 24 个时段列（表头为 "1".."24" 或整数列号）
    idx_hours = []
    for i, x in enumerate(h):
        try:
            if int(x) in range(1, 25):
                idx_hours.append(i)
        except Exception:
            pass
    # 兜底：连续 24 个数值列
    if len(idx_hours) != 24:
        idx_hours = list(range(idx_comp + 1, idx_comp + 1 + 24))[:24]
    for r in rows:
        if not r or r[idx_comp] is None:
            continue
        mk = month_key(r[idx_month])
        if mk is None:
            continue
        comp = str(r[idx_comp]).strip()
        if not comp:
            continue
        arr = []
        for i in idx_hours:
            v = r[i] if i < len(r) else None
            if v is None or v == "":
                arr.append(None)
            else:
                try:
                    arr.append(float(v))
                except Exception:
                    arr.append(None)
        price_map.setdefault(mk, {})[comp] = arr
    print(f"[price] {len(price_map)} months, {sum(len(v) for v in price_map.values())} companies")

    # ---- 电量 ----
    en_sheet = wb[wb.sheetnames[0]]
    h, rows = _read_header(en_sheet)
    time_cols = _detect_time_cols(h)
    if len(time_cols) < 50:
        raise SystemExit(f"电量 sheet 未识别到分时列，仅 {len(time_cols)} 个，实际 sheet={wb.sheetnames[0]}")
    TIME_LABELS = [t[3] for t in time_cols]
    name_col = next((i for i, x in enumerate(h) if "市场成员" in x), None)
    date_col = next((i for i, x in enumerate(h) if "日期" in x), None)
    meter_col = next((i for i, x in enumerate(h) if x == "计量点"), None)
    acc_col = next((i for i, x in enumerate(h) if x == "户号"), None)

    class CompData:
        __slots__ = ("slots", "days", "meters", "accs", "daily", "months", "max_load")
        def __init__(self):
            self.slots = [0.0] * len(time_cols)
            self.days = set()
            self.meters = set()
            self.accs = set()
            self.daily = {}
            self.months = {}
            self.max_load = 0.0

    companies = {}
    n_rows = 0
    for r in rows:
        n_rows += 1
        if n_rows % 20000 == 0:
            print(f"  row {n_rows} ...")
        if not r or name_col is None or r[name_col] is None:
            continue
        nm = str(r[name_col]).strip()
        if not nm:
            continue
        dl = r[date_col] if date_col is not None and r[date_col] is not None else None
        if dl is None:
            continue
        # 统一为 YYYY-MM-DD
        from datetime import datetime, timedelta
        if isinstance(dl, datetime):
            dl_str = dl.strftime("%Y-%m-%d")
        elif isinstance(dl, (int, float)) and dl > 20000:
            d = datetime(1899, 12, 30) + timedelta(days=int(dl))
            dl_str = d.strftime("%Y-%m-%d")
        else:
            dl_str = str(dl)[:10]
        mk = dl_str[:7]
        if nm not in companies:
            companies[nm] = CompData()
        c = companies[nm]
        c.days.add(dl_str)
        if meter_col is not None and r[meter_col] is not None:
            c.meters.add(r[meter_col])
        if acc_col is not None and r[acc_col] is not None:
            c.accs.add(r[acc_col])
        # 单次循环累加企业级和月级
        row_sum = 0.0
        if mk not in c.months:
            c.months[mk] = {"days": set(), "slotS": [0.0] * len(time_cols)}
        m = c.months[mk]
        m["days"].add(dl_str)
        for idx, (i, hh, mm, _lbl) in enumerate(time_cols):
            v = r[i] if i < len(r) else 0
            try:
                v = float(v) if v is not None else 0.0
            except Exception:
                v = 0.0
            c.slots[idx] += v
            m["slotS"][idx] += v
            row_sum += v
            if v > c.max_load:
                c.max_load = v
        c.daily[dl_str] = c.daily.get(dl_str, 0.0) + row_sum
    print(f"[read] {n_rows} rows, {len(companies)} companies, {time.time()-t0:.1f}s")
    wb.close()

    # ---- 组装 payload ----
    result_companies = {}
    all_months = set()
    total_overall = 0.0
    for nm, c in companies.items():
        n_days = max(1, len(c.days))
        curve = [round(v / n_days, 4) for v in c.slots]
        months_arr = []
        for mk in sorted(c.months.keys()):
            m = c.months[mk]
            nd = len(m["days"])
            slotS = m["slotS"]
            energy = sum(slotS)
            hourE = [0.0] * 24
            for idx, (_i, hh, mm, _lbl) in enumerate(time_cols):
                hourE[slot_hour(hh, mm)] += slotS[idx]
            pr = price_map.get(mk, {}).get(nm)
            cost = None
            avgp = None
            if pr and any(p is not None for p in pr):
                cost = sum(hourE[h] * pr[h] for h in range(24) if pr[h] is not None)
                avgp = cost / energy if energy else None
            months_arr.append({
                "m": mk, "days": nd, "energy": round(energy, 2),
                "mcurve": [round(v / nd, 4) for v in slotS],
                "hourE": [round(v, 4) for v in hourE],
                "price": pr,
                "cost": round(cost, 2) if cost is not None else None,
                "avgPrice": round(avgp, 2) if avgp is not None else None,
            })
            all_months.add(mk)
        months_arr.sort(key=lambda x: x["m"])
        dailylist = [{"d": d, "v": round(c.daily[d], 2)} for d in sorted(c.days)]
        result_companies[nm] = {
            "total": round(sum(c.slots), 2),
            "nMeter": len([m for m in c.meters if m is not None]),
            "nAccount": len([a for a in c.accs if a is not None]),
            "daily": dailylist,
            "curve": curve,
            "maxLoad": round(c.max_load, 4),
            "months": months_arr,
        }
        total_overall += sum(c.slots)

    MONTHS = sorted(all_months)
    names = list(result_companies.keys())
    overall = [round(sum(result_companies[n]["curve"][i] for n in names), 4) for i in range(len(time_cols))]
    payload = {
        "mode": "detailed",
        "timeCols": TIME_LABELS,
        "hourLabels": [f"{h}:00" for h in range(24)],
        "months": MONTHS,
        "overall": overall,
        "overallTotal": round(total_overall, 2),
        "companies": result_companies,
    }

# ================================================================
# 模式 B：结算汇总
# ================================================================
else:
    sheet = wb[wb.sheetnames[0]]
    h, rows = _read_header(sheet)
    def col_contains(substr):
        for i, x in enumerate(h):
            if substr in x:
                return i
        return None
    idx_month = col_contains("月")
    idx_comp = col_contains("企业") or col_contains("名称")
    idx_en = col_contains("结算电量") or col_contains("电量")
    idx_pr = col_contains("结算电价") or col_contains("电价")
    idx_plan = col_contains("方案")
    idx_contact = col_contains("联系人")
    if idx_month is None or idx_comp is None or idx_en is None:
        raise SystemExit(f"结算汇总表缺少必要列（month={idx_month}, comp={idx_comp}, energy={idx_en}, price={idx_pr}）。表头：{h}")
    print(f"[settlement cols] month={idx_month}(={h[idx_month]}), comp={idx_comp}(={h[idx_comp]}), energy={idx_en}(={h[idx_en]}), price={idx_pr}")

    companies = {}
    for r in rows:
        if not r or r[idx_comp] is None:
            continue
        nm = str(r[idx_comp]).strip()
        if not nm:
            continue
        mk = month_key(r[idx_month])
        if mk is None:
            continue
        en_v = 0.0
        try:
            en_v = float(r[idx_en]) if r[idx_en] is not None else 0.0
        except Exception:
            continue
        pr_v = None
        if idx_pr is not None:
            try:
                pr_v = float(r[idx_pr]) if r[idx_pr] is not None else None
            except Exception:
                pr_v = None
        plan = r[idx_plan] if idx_plan is not None else None
        contact = r[idx_contact] if idx_contact is not None else None
        if nm not in companies:
            companies[nm] = {"total": 0.0, "months": {}}
        c = companies[nm]
        if mk not in c["months"]:
            c["months"][mk] = {"energy": 0.0, "price": None, "cost": None, "plan": plan}
        m = c["months"][mk]
        m["energy"] += en_v
        if pr_v is not None:
            # 多个行时取最后一个非空值（同企业同月份理论上只有一行）
            m["price"] = pr_v
        c["total"] += en_v
    wb.close()

    result_companies = {}
    all_months = set()
    total_overall = 0.0
    for nm, c in companies.items():
        months_arr = []
        for mk in sorted(c["months"].keys()):
            m = c["months"][mk]
            cost = m["energy"] * m["price"] if m["price"] is not None else None
            months_arr.append({
                "m": mk,
                "energy": round(m["energy"], 2),
                "avgPrice": round(m["price"], 2) if m["price"] is not None else None,
                "cost": round(cost, 2) if cost is not None else None,
                "plan": m.get("plan"),
                # 分时字段置空（前端据此判断渲染简化视图）
                "mcurve": None, "hourE": None, "price": None, "days": None,
            })
            all_months.add(mk)
        months_arr.sort(key=lambda x: x["m"])
        total_energy = sum(c["months"][mk]["energy"] for mk in c["months"])
        # 综合电价 = 总电费 / 总电量
        total_cost = sum((c["months"][mk]["energy"] * c["months"][mk]["price"])
                         for mk in c["months"] if c["months"][mk]["price"] is not None)
        comp_avg_price = total_cost / total_energy if total_energy and total_cost else None
        result_companies[nm] = {
            "total": round(total_energy, 2),
            "totalCost": round(total_cost, 2) if total_cost else None,
            "avgPrice": round(comp_avg_price, 2) if comp_avg_price is not None else None,
            "months": months_arr,
            # 分时段字段置空，让前端自动切换到简化视图
            "daily": None, "curve": None, "maxLoad": None,
            "nMeter": 0, "nAccount": 0,
        }
        total_overall += total_energy

    MONTHS = sorted(all_months)
    payload = {
        "mode": "settlement",
        "months": MONTHS,
        "overall": None,
        "overallTotal": round(total_overall, 2),
        "companies": result_companies,
    }

with open(OUT, "w", encoding="utf-8") as f:
    f.write("window.DASH_DATA = ")
    json.dump(payload, f, ensure_ascii=False, separators=(",", ":"))
    f.write(";\n")
print(f"[done] data.js = {os.path.getsize(OUT)/1024:.1f} KB, "
      f"mode={mode}, companies={len(result_companies)}, months={MONTHS}, total={round(total_overall,1)} MWh, "
      f"elapsed {time.time()-t0:.1f}s")